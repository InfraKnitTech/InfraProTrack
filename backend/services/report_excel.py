from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable


def build_excel_workbook(
    title: str,
    sheets: Iterable[tuple[str, list[str], list[list[Any]]]],
    metadata: list[tuple[str, str]] | None = None,
) -> BytesIO:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise RuntimeError("Excel export requires openpyxl. Install backend requirements first.") from exc

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    if metadata:
        sheet = workbook.create_sheet("Report Info")
        sheet.append([title])
        sheet["A1"].font = Font(bold=True, size=14)
        sheet.append([])
        for key, value in metadata:
            sheet.append([key, value])
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 42

    header_fill = PatternFill(fill_type="solid", fgColor="1D4ED8")
    header_font = Font(color="FFFFFF", bold=True)
    subtitle_fill = PatternFill(fill_type="solid", fgColor="DBEAFE")

    for sheet_name, headers, rows in sheets:
        sheet = workbook.create_sheet(sheet_name[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in rows:
            sheet.append(row)

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, start=1):
            values = [header]
            for row in rows[:250]:
                if index - 1 < len(row):
                    values.append("" if row[index - 1] is None else str(row[index - 1]))
            max_len = max((len(str(value)) for value in values), default=10)
            sheet.column_dimensions[get_column_letter(index)].width = min(max(max_len + 2, 12), 42)

        if not rows:
            sheet.append(["No rows available"] + [""] * max(0, len(headers) - 1))
            for cell in sheet[2]:
                cell.fill = subtitle_fill

    if not workbook.sheetnames:
        workbook.create_sheet("Report")

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
